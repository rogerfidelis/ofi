from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


# ============================================================
# CONFIGURAÇÕES
# ============================================================

BASE = Path(__file__).resolve().parents[1]

EMPRESAS = ["cbo", "bram", "starnav"]

ARQUIVO_ORIGEM = (
    BASE
    / "dados"
    / "SURVEY_VESSELS"
    / "SURVEY_VESSELS_20260914_tarde.xlsx"
)

# Coluna do arquivo de origem -> coluna do arquivo de destino
MAPEAMENTO_COLUNAS = {
    "data": "data_consulta",
    "lat": "lat_a",
    "lon": "lon_a",
    "data reportada": "data_reportada",
    "status": "status",
}

# Todas as colunas gravadas participam da identificação
# de duplicidade.
CHAVE_DUPLICIDADE = [
    "data_consulta",
    "lat_a",
    "lon_a",
    "data_reportada",
    "status",
]


# ============================================================
# NORMALIZAÇÃO
# ============================================================

def valor_vazio(valor):
    """Verifica valores vazios com segurança."""
    if valor is None:
        return True

    try:
        return bool(pd.isna(valor))
    except (TypeError, ValueError):
        return False


def normalizar_texto(valor):
    if valor_vazio(valor):
        return None

    texto = str(valor).strip()

    if not texto:
        return None

    return texto.upper()


def normalizar_cabecalho(valor):
    if valor_vazio(valor):
        return ""

    return str(valor).strip().lower()


def normalizar_nome_aba(valor):
    if valor_vazio(valor):
        return ""

    return str(valor).strip().upper()


def normalizar_data(valor):
    """
    Converte datas de diferentes formatos para YYYY-MM-DD.
    Isso permite comparar datas do pandas com datas do Excel.
    """
    if valor_vazio(valor):
        return None

    if isinstance(valor, pd.Timestamp):
        return valor.date().isoformat()

    if isinstance(valor, datetime):
        return valor.date().isoformat()

    if isinstance(valor, date):
        return valor.isoformat()

    convertido = pd.to_datetime(
        valor,
        dayfirst=True,
        errors="coerce",
    )

    if pd.isna(convertido):
        return str(valor).strip()

    return convertido.date().isoformat()


def normalizar_coordenada(valor):
    """
    Arredonda a coordenada para seis casas decimais,
    evitando diferenças causadas apenas por formatação.
    """
    if valor_vazio(valor):
        return None

    try:
        return round(float(valor), 6)
    except (TypeError, ValueError):
        return str(valor).strip()


def normalizar_valor_chave(nome_coluna, valor):
    if nome_coluna in {"data_consulta", "data_reportada"}:
        return normalizar_data(valor)

    if nome_coluna in {"lat_a", "lon_a"}:
        return normalizar_coordenada(valor)

    if nome_coluna == "status":
        return normalizar_texto(valor)

    return valor


def converter_valor_excel(valor):
    """Converte valores do pandas para formatos aceitos pelo Excel."""
    if valor_vazio(valor):
        return None

    if isinstance(valor, pd.Timestamp):
        return valor.to_pydatetime()

    return valor


# ============================================================
# CABEÇALHOS
# ============================================================

def obter_cabecalhos_dataframe(dataframe):
    """
    Exemplo de resultado:
    {
        "embarcacao": "Embarcacao",
        "data": "Data",
        "lat": "Lat"
    }
    """
    return {
        normalizar_cabecalho(coluna): coluna
        for coluna in dataframe.columns
    }


def obter_cabecalhos_planilha(ws):
    """
    Retorna:
    cabeçalho normalizado -> número da coluna no Excel.
    """
    cabecalhos = {}

    for numero_coluna in range(1, ws.max_column + 1):
        valor = ws.cell(
            row=1,
            column=numero_coluna,
        ).value

        nome = normalizar_cabecalho(valor)

        if nome:
            cabecalhos[nome] = numero_coluna

    return cabecalhos


# ============================================================
# DUPLICIDADE
# ============================================================

def criar_chave_registro(registro):
    """
    Cria uma chave comparável com os cinco campos gravados.

    Exemplo:
    (
        "2026-08-03",
        -22.912345,
        -43.123456,
        "2026-08-03",
        "UNDERWAY"
    )
    """
    return tuple(
        normalizar_valor_chave(
            nome_coluna,
            registro.get(nome_coluna),
        )
        for nome_coluna in CHAVE_DUPLICIDADE
    )


def carregar_chaves_existentes(ws, cabecalhos_destino):
    """
    Lê os registros existentes na aba da embarcação.
    """
    chaves_existentes = set()

    for numero_linha in range(2, ws.max_row + 1):
        registro = {
            nome_coluna: ws.cell(
                row=numero_linha,
                column=cabecalhos_destino[nome_coluna],
            ).value
            for nome_coluna in CHAVE_DUPLICIDADE
        }

        chave = criar_chave_registro(registro)

        # Não considera linhas completamente vazias.
        if any(valor is not None for valor in chave):
            chaves_existentes.add(chave)

    return chaves_existentes


# ============================================================
# PROCESSAMENTO POR EMPRESA
# ============================================================

def processar_empresa(empresa):
    arquivo_destino = (
        BASE
        / "dados"
        / f"SURVEY_{empresa.upper()}.xlsx"
    )

    print()
    print("=" * 70)
    print(f"EMPRESA: {empresa.upper()}")
    print("=" * 70)

    if not ARQUIVO_ORIGEM.exists():
        raise FileNotFoundError(
            f"Arquivo de origem não encontrado: {ARQUIVO_ORIGEM}"
        )

    if not arquivo_destino.exists():
        raise FileNotFoundError(
            f"Arquivo de destino não encontrado: {arquivo_destino}"
        )

    origem = pd.read_excel(
        ARQUIVO_ORIGEM,
        sheet_name=empresa,
    )

    if origem.empty:
        print(f"A aba '{empresa}' está vazia.")
        return

    cabecalhos_origem = obter_cabecalhos_dataframe(origem)

    colunas_origem_obrigatorias = {
        "embarcacao",
        "data",
        "lat",
        "lon",
        "data reportada",
        "status",
    }

    colunas_origem_ausentes = (
        colunas_origem_obrigatorias
        - set(cabecalhos_origem)
    )

    if colunas_origem_ausentes:
        raise ValueError(
            "Colunas ausentes no arquivo de origem: "
            + ", ".join(sorted(colunas_origem_ausentes))
        )

    wb = load_workbook(arquivo_destino)

    # Evita reler uma aba toda vez que a embarcação aparece.
    cache_abas = {}

    inseridos = 0
    duplicados = 0
    invalidos = 0
    abas_nao_encontradas = set()
    abas_com_erro = set()

    coluna_embarcacao = cabecalhos_origem["embarcacao"]

    for indice, linha in origem.iterrows():
        numero_linha_origem = indice + 2

        aba_destino = normalizar_nome_aba(
            linha[coluna_embarcacao]
        )

        if not aba_destino:
            invalidos += 1

            print(
                f"Linha {numero_linha_origem}: "
                "embarcação não informada."
            )

            continue

        if aba_destino not in wb.sheetnames:
            abas_nao_encontradas.add(aba_destino)

            print(
                f"Linha {numero_linha_origem}: "
                f"aba '{aba_destino}' não encontrada."
            )

            continue

        # Prepara e armazena as informações da aba.
        if aba_destino not in cache_abas:
            ws = wb[aba_destino]
            cabecalhos_destino = obter_cabecalhos_planilha(ws)

            colunas_destino_obrigatorias = set(
                MAPEAMENTO_COLUNAS.values()
            )

            colunas_destino_ausentes = (
                colunas_destino_obrigatorias
                - set(cabecalhos_destino)
            )

            if colunas_destino_ausentes:
                abas_com_erro.add(aba_destino)

                print(
                    f"Aba '{aba_destino}' sem os cabeçalhos: "
                    + ", ".join(
                        sorted(colunas_destino_ausentes)
                    )
                )

                continue

            cache_abas[aba_destino] = {
                "ws": ws,
                "cabecalhos": cabecalhos_destino,
                "chaves": carregar_chaves_existentes(
                    ws,
                    cabecalhos_destino,
                ),
            }

        if aba_destino in abas_com_erro:
            continue

        dados_aba = cache_abas[aba_destino]

        ws = dados_aba["ws"]
        cabecalhos_destino = dados_aba["cabecalhos"]
        chaves_existentes = dados_aba["chaves"]

        # Monta o registro usando os nomes do arquivo destino.
        registro = {}

        for coluna_origem, coluna_destino in MAPEAMENTO_COLUNAS.items():
            nome_original = cabecalhos_origem[coluna_origem]
            registro[coluna_destino] = linha[nome_original]

        chave = criar_chave_registro(registro)

        # Evita inserir uma linha sem informações.
        if all(valor is None for valor in chave):
            invalidos += 1

            print(
                f"Linha {numero_linha_origem}: "
                "registro vazio — ignorado."
            )

            continue

        if chave in chaves_existentes:
            duplicados += 1

            print(
                f"Linha {numero_linha_origem}: "
                f"registro de '{aba_destino}' já existe — ignorado."
            )

            continue

        linha_destino = ws.max_row + 1 #ADICIONAR DADO NA LINHA ABAIXO

        for nome_coluna, valor in registro.items():
            numero_coluna = cabecalhos_destino[nome_coluna]

            ws.cell(
                row=linha_destino,
                column=numero_coluna,
                value=converter_valor_excel(valor),
            )

        # Registra imediatamente a chave para impedir duplicações
        # encontradas posteriormente na mesma execução.
        chaves_existentes.add(chave)

        inseridos += 1

        print(
            f"Linha {numero_linha_origem}: "
            f"'{aba_destino}' inserida na linha {linha_destino}."
        )

    wb.save(arquivo_destino)
    wb.close()

    print()
    print(f"Arquivo: {arquivo_destino.name}")
    print(f"Novos registros: {inseridos}")
    print(f"Duplicados ignorados: {duplicados}")
    print(f"Registros inválidos: {invalidos}")
    print(f"Abas não encontradas: {len(abas_nao_encontradas)}")

    if abas_nao_encontradas:
        print(
            "Embarcações sem aba: "
            + ", ".join(sorted(abas_nao_encontradas))
        )


# ============================================================
# EXECUÇÃO
# ============================================================

def main():
    inicio = datetime.now()
    empresas_com_erro = []

    print("=" * 70)
    print("COMBINADOR DE DADOS — SURVEY VESSELS")
    print("=" * 70)
    print(f"Arquivo de origem: {ARQUIVO_ORIGEM}")

    for empresa in EMPRESAS:
        try:
            processar_empresa(empresa)

        except Exception as erro:
            empresas_com_erro.append(empresa.upper())

            print()
            print(
                f"Erro ao processar {empresa.upper()}: {erro}"
            )

    duracao = datetime.now() - inicio

    print()
    print("=" * 70)
    print("PROCESSAMENTO FINALIZADO")
    print("=" * 70)
    print(f"Duração: {duracao}")

    if empresas_com_erro:
        print(
            "Empresas com erro: "
            + ", ".join(empresas_com_erro)
        )
    else:
        print("Todas as empresas foram processadas com sucesso.")


if __name__ == "__main__":
    main()